"""
Onboarding (docs/plan-global-pool.md, phase 2): one form, then a background build.

The build reads everything the user gave (CVs, notes, GitHub, portfolio), keeps only the lines
those documents back (app/provenance.py), works out years of experience, writes one baseline
resume per target role family and renders each to one page. One review screen follows. Nothing
is asked along the way, and a line that cannot be backed is left out rather than questioned.
"""
from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Optional

import httpx
from docx import Document
from psycopg.types.json import Jsonb
from pypdf import PdfReader

from . import llm
from .db import system_tx, user_tx
from .pipeline import prompts
from .pipeline.schemas import (
    BaselineSet, ExEducation, ExEntry, ExItem, ExProfile, Extraction, ProposedBaseline,
)
from .provenance import (
    Check, Corpus, check_line, check_name, drop_unbacked_sentences, has_contact, has_url,
    keep_if_numbers_known, keep_known_tools, numbers, url_key,
)
from .taxonomy import (
    FAMILIES, band_for_years, canonical_skill, family_with_neighbours, is_internship, skill_key,
    title_family, years_of_experience,
)

MAX_FAMILIES = 4
README_REPOS = 6
README_CHARS = 4000
FIT_RANK = {"strong": 0, "good": 1, "stretch": 2}
SENIOR_WORDS = re.compile(r"\b(senior|sr\.?|lead|staff|principal|head)\s+", re.I)
GITHUB_URL_RE = re.compile(r"^(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9-]{1,39})/?(?:[?#].*)?$", re.I)
GITHUB_NAME_RE = re.compile(r"^@?([A-Za-z0-9](?:[A-Za-z0-9-]{0,38}))$")
SECTION_DEFAULTS = {"awards": ("Awards & Certifications", "list", 0),
                    "roles": ("Other Roles & Responsibilities", "detail", 1)}
NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")


# ------------------------------------------------------------------ reading what the user gave

def text_from_upload(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        # Two-column PDFs interleave columns; the provenance check searches wide windows.
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    parts.extend(p.text for p in cell.paragraphs)
        return "\n".join(p for p in parts if p.strip())
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", "replace")
    raise ValueError("Upload a PDF, DOCX or TXT file")


def github_username(s: str) -> Optional[str]:
    s = (s or "").strip()
    m = GITHUB_URL_RE.match(s) or GITHUB_NAME_RE.match(s)
    return m.group(1) if m else None


def _github_headers(raw: bool = False) -> dict:
    h = {"Accept": "application/vnd.github.raw+json" if raw else "application/vnd.github+json",
         "User-Agent": "jobreach"}
    if os.environ.get("GITHUB_TOKEN"):
        h["Authorization"] = "Bearer " + os.environ["GITHUB_TOKEN"]
    return h


def github_repos(username: str, timeout: float = 10.0) -> Optional[list[dict]]:
    """The user's own public repositories, most recently pushed first. Forks and archived repos
    are left out. Public API, no key (GITHUB_TOKEN raises the rate limit)."""
    try:
        r = httpx.get("https://api.github.com/users/{}/repos".format(username), headers=_github_headers(),
                      params={"per_page": 100, "sort": "pushed", "type": "owner"}, timeout=timeout)
    except httpx.HTTPError:
        return None
    if r.status_code != 200:
        return None
    return [x for x in r.json() if not x.get("fork") and not x.get("archived")][:30]


def github_summary(username: str, repos: list[dict]) -> str:
    lines = ["GitHub profile https://github.com/{}: {} public repositories of their own, most "
             "recently pushed first.".format(username, len(repos))]
    for x in repos:
        parts = [x["name"]]
        if x.get("description"):
            parts.append(x["description"])
        if x.get("language"):
            parts.append("language: " + x["language"])
        if x.get("topics"):
            parts.append("topics: " + ", ".join(x["topics"]))
        if x.get("homepage"):
            parts.append("live: " + x["homepage"])
        if x.get("stargazers_count"):
            parts.append("stars: {}".format(x["stargazers_count"]))
        parts.append("repo: " + x["html_url"])
        parts.append("last push: " + (x.get("pushed_at") or "")[:10])
        lines.append("- " + " | ".join(parts))
    return "\n".join(lines)


def github_readmes(username: str, repos: list[dict], timeout: float = 10.0) -> list[tuple[str, str]]:
    """The READMEs of the most recent repositories: where students describe projects their CV
    leaves out. With no interview any more, this is the main second source."""
    out = []
    for x in repos[:README_REPOS]:
        try:
            r = httpx.get("https://api.github.com/repos/{}/{}/readme".format(username, x["name"]),
                          headers=_github_headers(raw=True), timeout=timeout)
        except httpx.HTTPError:
            continue
        if r.status_code == 200 and r.text.strip():
            text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", r.text)          # badges and images
            out.append((x["name"], text[:README_CHARS]))
    return out


def portfolio_text(url: str) -> Optional[str]:
    """The user's own portfolio page. LinkedIn is never fetched (spec 11.1)."""
    from .pipeline.verify import fetch_page

    u = url.strip()
    if not u:
        return None
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    if "linkedin.com" in u.lower():
        return None
    return fetch_page(u)


# ------------------------------------------------------------------ the form

@dataclass
class Form:
    stage: str
    families: list[str]
    titles: list[str] = field(default_factory=list)
    github: str = ""
    portfolio: str = ""
    linkedin: str = ""
    other_links: str = ""
    notes: str = ""
    open_to: list[str] = field(default_factory=list)
    remote_ok: bool = True
    hybrid_ok: bool = True
    onsite_ok: bool = True
    locations: list[str] = field(default_factory=list)
    stipend_floor: Optional[int] = None
    salary_floor: Optional[int] = None
    notice_period: Optional[str] = None
    start_date: Optional[str] = None
    excluded_companies: list[str] = field(default_factory=list)
    skip_big_tech: Optional[bool] = None


def families_for(chosen: list[str], titles: list[str]) -> list[str]:
    """The target role families: the ones picked, then those the typed titles name. At most
    four, since each is a resume to keep up."""
    out: list[str] = []
    for f in list(chosen) + [title_family(t) for t in titles]:
        if f in FAMILIES and f != "other" and f not in out:
            out.append(f)
    return out[:MAX_FAMILIES]


def role_types(families: list[str]) -> list[str]:
    """What the per-lead screen accepts: the targets and their neighbours."""
    return sorted({n for f in families for n in family_with_neighbours(f)})


def stage_defaults(stage: str) -> dict:
    """Students look for internships and skip big tech by default (the reference data);
    graduates and experienced people want full-time work anywhere."""
    if stage == "student":
        return {"open_to": ["internship"], "excluded_company_types": ["big_tech"]}
    return {"open_to": ["full_time"], "excluded_company_types": []}


def header_links(form: Form) -> list[dict]:
    links = []
    name = github_username(form.github)
    if name:
        links.append({"text": "github.com/" + name, "url": "https://github.com/" + name})
    for raw in (form.linkedin, form.portfolio):
        u = raw.strip()
        if u:
            full = u if re.match(r"^https?://", u, re.I) else "https://" + u
            links.append({"text": url_key(full), "url": full})
    return links


def apply_form(conn, user_id: str, form: Form, uploads: list[tuple[str, str]]) -> None:
    """Store the form: the documents, the profile, preferences with stage defaults, and a
    queued build. Call inside user_tx(user_id); the caller enqueues the build."""
    conn.execute("delete from source_documents")
    for name, text in uploads:
        conn.execute("insert into source_documents (user_id, kind, filename, text) values (%s, 'resume', %s, %s)",
                     (user_id, name, text))
    if form.notes.strip():
        conn.execute("insert into source_documents (user_id, kind, filename, text) values (%s, 'about', 'your notes', %s)",
                     (user_id, form.notes.strip()))
    if form.other_links.strip():
        conn.execute("insert into source_documents (user_id, kind, filename, text) values (%s, 'links', 'links', %s)",
                     (user_id, form.other_links.strip()))

    gh = github_username(form.github)
    conn.execute(
        """update profiles set career_stage = %s, about = coalesce(nullif(%s, ''), about),
               github_url = %s, portfolio_url = nullif(%s, ''), linkedin_url = nullif(%s, ''),
               links = %s, onboarding_step = 'building',
               build = jsonb_build_object('status', 'queued', 'step', 'reading', 'started_at', now()),
               updated_at = now()""",
        (form.stage, form.notes.strip(), "https://github.com/" + gh if gh else None, form.portfolio.strip(),
         form.linkedin.strip(), Jsonb(header_links(form))))

    families = families_for(form.families, form.titles)
    d = stage_defaults(form.stage)
    excluded_types = (["big_tech"] if form.skip_big_tech else []) if form.skip_big_tech is not None \
        else d["excluded_company_types"]
    conn.execute(
        """update preferences set target_families = %s, role_types = %s, desired_roles = %s, target_roles = %s,
               open_to = %s, remote_ok = %s, hybrid_ok = %s, onsite_ok = %s, locations = %s,
               stipend_floor = %s, salary_floor = %s, notice_period = %s,
               start_date = coalesce(nullif(%s, ''), start_date), excluded_companies = %s,
               excluded_company_types = %s, updated_at = now()""",
        (families, role_types(families), form.titles, form.titles, form.open_to or d["open_to"],
         form.remote_ok, form.hybrid_ok, form.onsite_ok, form.locations, form.stipend_floor,
         form.salary_floor, form.notice_period, form.start_date or "", form.excluded_companies,
         excluded_types))


# ------------------------------------------------------------------ checking the extraction (pure)

def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:40] or "item"


@dataclass
class Line:
    text: str
    check: Check


@dataclass
class CheckedItem:
    item: ExItem
    check: Check
    bullets: list[Line]
    trimmed: list[str] = field(default_factory=list)

    @property
    def usable(self) -> bool:
        return self.check.ok


@dataclass
class CheckedProfile:
    profile: ExProfile
    items: list[CheckedItem]
    education: list[tuple[ExEducation, Check]]
    entries: list[tuple[ExEntry, Check]]
    skills: list[Line]
    other_facts: list[Line]

    def counts(self) -> dict:
        lines = [b.check for it in self.items for b in it.bullets] + [c for _, c in self.entries] \
            + [s.check for s in self.skills] + [c for _, c in self.education] + [it.check for it in self.items]
        return {"kept": sum(c.ok for c in lines), "left_out": sum(not c.ok for c in lines)}


def check_extraction(ex: Extraction, corpus: Corpus) -> CheckedProfile:
    """Hold every extracted line to the documents. Headers are trimmed (a tool, a period or a
    tagline the documents don't back is dropped from the header); whole lines that fail are
    kept as left out, with the reason."""
    p = ex.profile
    profile = p.model_copy(update={
        "name": p.name if p.name and check_name(p.name, corpus).ok else None,
        "email": p.email if has_contact(p.email, corpus) else None,
        "phone": p.phone if has_contact(p.phone, corpus) else None,
        "location": p.location if p.location and check_line(p.location, corpus).ok else None,
        "headline": None,                                   # baselines write their own title lines
        "grad_date": keep_if_numbers_known(p.grad_date, corpus),
        "batch_year": p.batch_year if p.batch_year and str(p.batch_year) in corpus.nums else None,
        "cgpa": p.cgpa if p.cgpa is not None and numbers(str(p.cgpa)) <= corpus.nums else None,
        "links": [lk for lk in p.links if has_url(lk.url, corpus)],
    })

    items, seen = [], set()
    for it in ex.items:
        name_check = check_name(it.name, corpus)
        trimmed = []
        tagline = it.tagline
        if tagline and not check_line(tagline, corpus).ok:
            trimmed.append("tagline '{}'".format(tagline))
            tagline = None
        period = keep_if_numbers_known(it.period, corpus)
        if it.period and not period:
            trimmed.append("period '{}'".format(it.period))
        stack, dropped = keep_known_tools(it.stack, corpus)
        trimmed += ["'{}' from the stack".format(d) for d in dropped]
        links = [lk for lk in it.links if has_url(lk.url, corpus)]
        bullets, texts = [], set()
        for b in it.bullets:
            k = " ".join(b.lower().split())
            if k in texts:
                continue
            texts.add(k)
            bullets.append(Line(b, check_line(b, corpus) if name_check.ok
                                else Check(False, "Its project or job wasn't found in what you gave us")))
        key = _slug(it.key or it.name)
        while key in seen:
            key += "_2"
        seen.add(key)
        items.append(CheckedItem(it.model_copy(update={"key": key, "tagline": tagline, "period": period,
                                                       "stack": stack, "links": links}),
                                 name_check, bullets, trimmed))

    education = []
    for ed in ex.education:
        c = check_line(ed.institution, corpus, coverage=0.8)
        education.append((ed.model_copy(update={
            "degree": ed.degree if ed.degree and check_line(ed.degree, corpus).ok else None,
            "meta": keep_if_numbers_known(ed.meta, corpus),
            "result": keep_if_numbers_known(ed.result, corpus),
            "lines": [ln for ln in ed.lines if check_line(ln, corpus).ok],
        }), c))

    entries = [(e, check_line(((e.lead or "") + " " + e.text).strip(), corpus)) for e in ex.entries]

    skills, keys = [], set()
    for s in ex.skills:
        name = canonical_skill(s)
        k = skill_key(name)
        if not k or k in keys:
            continue
        keys.add(k)
        ok = corpus.has_tech(s)
        skills.append(Line(name, Check(True) if ok else Check(False, "{} is not mentioned in anything you gave us".format(name))))

    other = [Line(t, check_line(t, corpus)) for t in ex.other_facts]
    return CheckedProfile(profile, items, education, entries, skills, other)


def counted_experience(items: list[CheckedItem], today: Optional[date] = None) -> float:
    """Full-time years: usable jobs that are full-time, or untyped and not called an internship."""
    periods = [ci.item.period for ci in items
               if ci.usable and ci.item.kind == "experience"
               and (ci.item.employment == "full_time"
                    or (ci.item.employment == "unknown"
                        and not is_internship("{} {}".format(ci.item.name, ci.item.tagline or ""))))]
    return years_of_experience(periods, today)


# ------------------------------------------------------------------ baselines (pure)

def fallback_title(family: str) -> str:
    f = FAMILIES.get(family)
    return f.aliases[0] if f and f.aliases else "Software Engineer"


def _cap_fit(fit: str, evidence: list[str]) -> str:
    cap = "strong" if len(evidence) >= 2 else "good" if evidence else "stretch"
    return fit if FIT_RANK.get(fit, 2) >= FIT_RANK[cap] else cap


def validate_baselines(out: BaselineSet, families: list[str], items: list[dict], skills: list[str],
                       corpus: Corpus, band: str, extra_numbers: set[str]) -> tuple[list[dict], list[dict]]:
    """Hold the model's plan to the record: only usable items and listed skills, a fit no
    higher than the evidence supports, no seniority the band does not show, and no sentence
    carrying a number or tool the documents lack. A family the model skipped still gets a
    plain baseline. Returns (track rows, suggestions)."""
    item_keys = [i["key"] for i in items]
    known = set(item_keys)
    experience_first = [i["key"] for i in items if i["kind"] == "experience"] + \
        [i["key"] for i in items if i["kind"] == "project"]
    skill_by_key = {skill_key(canonical_skill(s)): s for s in skills}
    allow_senior = band in ("senior", "lead")
    by_family: dict[str, ProposedBaseline] = {}
    for b in out.baselines:
        if b.family in families and b.family not in by_family:
            by_family[b.family] = b

    tracks = []
    for i, fam in enumerate(families):
        b = by_family.get(fam)
        label = FAMILIES[fam].label
        if b is None:
            tracks.append({"key": fam, "label": label, "role_family": fam, "title_line": fallback_title(fam),
                           "summary": "", "left_sections": _default_sections(items, experience_first),
                           "skills": [{"label": "Skills", "items": ", ".join(skills[:16])}] if skills else [],
                           "fit": "stretch", "fit_why": "", "gaps": [], "sort": i})
            continue
        left = [{"heading": s.heading.strip() or "Projects", "item_keys": [k for k in dict.fromkeys(s.item_keys) if k in known]}
                for s in b.left_sections]
        left = [s for s in left if s["item_keys"]] or _default_sections(items, experience_first)
        groups = []
        for g in b.skills:
            kept = [skill_by_key[k] for k in (skill_key(canonical_skill(x)) for x in g.items.split(","))
                    if k in skill_by_key]
            if kept:
                groups.append({"label": g.label.strip() or "Skills", "items": ", ".join(dict.fromkeys(kept))})
        title = b.title_line.strip() or fallback_title(fam)
        if not allow_senior:
            title = SENIOR_WORDS.sub("", title)
        if numbers(title) - corpus.nums - extra_numbers:
            title = fallback_title(fam)
        summary, _ = drop_unbacked_sentences(b.summary, corpus, extra_numbers)
        evidence = [k for k in dict.fromkeys(b.evidence_item_keys) if k in known]
        tracks.append({"key": fam, "label": label, "role_family": fam, "title_line": title,
                       "summary": summary, "left_sections": left, "skills": groups,
                       "fit": _cap_fit(b.fit, evidence), "fit_why": b.fit_why.strip(),
                       "gaps": [g.strip() for g in b.gaps if g.strip()][:4], "sort": i})

    suggestions = []
    for s in out.suggestions:
        evidence = [k for k in dict.fromkeys(s.evidence_item_keys) if k in known]
        if s.family in FAMILIES and s.family not in families and s.family != "other" and len(evidence) >= 2 \
                and s.family not in {x["family"] for x in suggestions}:
            suggestions.append({"family": s.family, "label": FAMILIES[s.family].label, "why": s.why.strip()})
    return tracks, suggestions[:2]


def _default_sections(items: list[dict], order: list[str]) -> list[dict]:
    kinds = {i["key"]: i["kind"] for i in items}
    exp = [k for k in order if kinds[k] == "experience"][:2]
    proj = [k for k in order if kinds[k] == "project"][:4 - len(exp)]
    out = []
    if exp:
        out.append({"heading": "Experience", "item_keys": exp})
    if proj:
        out.append({"heading": "Projects", "item_keys": proj})
    return out


def baseline_context(families: list[str], person: dict, items: list[dict], skills: list[str],
                     entries: list[str]) -> str:
    return llm.dumps({
        "targets": [{"family": f, "label": FAMILIES[f].label} for f in families],
        "person": person, "items": items, "skills": skills, "awards_and_roles": entries,
    })


def extra_numbers_for(person: dict) -> set[str]:
    """Numbers a title or summary may carry beyond the documents: the computed years."""
    y = person.get("experience_years")
    out = set()
    if y:
        out |= {str(int(y)), "{:g}".format(float(y))}
    return out


# ------------------------------------------------------------------ the build (database + model)

def set_build(user_id: str, **fields) -> None:
    with user_tx(user_id) as conn:
        conn.execute("update profiles set build = build || %s::jsonb", (Jsonb(fields),))


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _section_id(conn, user_id: str, key: str) -> str:
    heading, style, sort = SECTION_DEFAULTS[key]
    row = conn.execute(
        """insert into resume_sections (user_id, key, heading, style, sort) values (%s, %s, %s, %s, %s)
           on conflict (user_id, key) do update set key = excluded.key returning id::text""",
        (user_id, key, heading, style, sort)).fetchone()
    return row["id"]


def _gather_sources(user_id: str) -> list[tuple[str, str]]:
    """The documents as stored, plus GitHub (repositories and READMEs) and the portfolio,
    fetched now. Returns (label, text) pairs for the corpus and the extraction."""
    with user_tx(user_id) as conn:
        prof = conn.execute("select github_url, portfolio_url from profiles").fetchone() or {}
    fetched_docs, status = [], {}
    name = github_username(prof.get("github_url") or "")
    if name:
        repos = github_repos(name)
        status["github"] = {"username": name, "read": repos is not None}
        if repos:
            fetched_docs.append(("GitHub repositories", github_summary(name, repos)))
            fetched_docs += [("GitHub README: " + n, t) for n, t in github_readmes(name, repos)]
    if prof.get("portfolio_url"):
        text = portfolio_text(prof["portfolio_url"])
        status["portfolio"] = {"url": prof["portfolio_url"], "read": text is not None}
        if text:
            fetched_docs.append(("Portfolio", text))
    with user_tx(user_id) as conn:
        # Kept with the uploads, so later steps (adding a role family) search the same corpus.
        conn.execute("delete from source_documents where filename like 'fetched: %%'")
        for label, text in fetched_docs:
            conn.execute("insert into source_documents (user_id, kind, filename, text) values (%s, 'links', %s, %s)",
                         (user_id, "fetched: " + label, text))
        docs = conn.execute("select kind, filename, text from source_documents order by created_at").fetchall()
    set_build(user_id, **status)
    return [(source_label(d["kind"], d["filename"]), d["text"]) for d in docs]


def source_label(kind: str, filename: Optional[str]) -> str:
    name = (filename or kind).removeprefix("fetched: ")
    return "CV: " + name if kind == "resume" else name


def stored_corpus(conn) -> Corpus:
    return Corpus([(source_label(d["kind"], d["filename"]), d["text"]) for d in
                   conn.execute("select kind, filename, text from source_documents order by created_at").fetchall()])


def extract_profile(docs: list[tuple[str, str]], ctx: llm.CallContext) -> Extraction:
    volatile = "\n\n".join('<document name="{}">\n{}\n</document>'.format(label, text[:30000]) for label, text in docs)
    return llm.structured("onb_extract", Extraction, stable=[prompts.ONBOARD_EXTRACT], volatile=volatile,
                          effort="medium", max_tokens=16000, ctx=ctx)


def store_checked(conn, user_id: str, checked: CheckedProfile) -> None:
    """Replace the fact bank with the checked extraction. Call inside user_tx(user_id)."""
    for t in ("items", "entries", "education", "facts"):
        conn.execute("delete from {}".format(t))
    p = checked.profile
    conn.execute(
        """update profiles set name = coalesce(%s, name), email = coalesce(%s, email), phone = coalesce(%s, phone),
               location = coalesce(%s, location), grad_date = coalesce(%s, grad_date),
               batch_year = coalesce(%s, batch_year), cgpa = coalesce(%s, cgpa), updated_at = now()""",
        (p.name, p.email, p.phone, p.location, p.grad_date, p.batch_year, p.cgpa))
    if p.links:
        cur = conn.execute("select links from profiles").fetchone()["links"] or []
        have = {url_key(lk["url"]) for lk in cur}
        extra = [lk.model_dump() for lk in p.links if url_key(lk.url) not in have]
        conn.execute("update profiles set links = %s", (Jsonb(cur + extra),))

    for i, ci in enumerate(checked.items):
        it = ci.item
        prov = {**ci.check.as_json(), "trimmed": ci.trimmed}
        row = conn.execute(
            """insert into items (user_id, key, kind, name, tagline, period, stack, links, sort, confirmed, provenance)
               values (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) returning id""",
            (user_id, it.key, it.kind, it.name, it.tagline, it.period, it.stack,
             Jsonb([lk.model_dump() for lk in it.links]), 100 + i, ci.usable, Jsonb(prov))).fetchone()
        for j, b in enumerate(ci.bullets):
            f = conn.execute(
                """insert into facts (user_id, kind, text, source, confirmed_at, provenance)
                   values (%s, %s, %s, 'upload', case when %s then now() end, %s) returning id""",
                (user_id, "role" if it.kind == "experience" else "project", b.text, b.check.ok,
                 Jsonb(b.check.as_json()))).fetchone()
            conn.execute(
                """insert into bullets (user_id, item_id, text, fact_ids, has_metric, confirmed, sort, provenance)
                   values (%s, %s, %s, %s, %s, %s, %s, %s)""",
                (user_id, row["id"], b.text, [f["id"]], bool(NUMBER_RE.search(b.text)), b.check.ok, j,
                 Jsonb(b.check.as_json())))

    for i, (e, c) in enumerate(checked.entries):
        sid = _section_id(conn, user_id, e.section)
        f = conn.execute(
            """insert into facts (user_id, kind, text, source, confirmed_at, provenance)
               values (%s, %s, %s, 'upload', case when %s then now() end, %s) returning id""",
            (user_id, "award" if e.section == "awards" else "role", ((e.lead or "") + " " + e.text).strip(),
             c.ok, Jsonb(c.as_json()))).fetchone()
        text = e.text if not e.lead or e.text.startswith((" ", "-")) else " - " + e.text
        conn.execute(
            """insert into entries (user_id, section_id, lead, text, fact_ids, confirmed, sort, provenance)
               values (%s, %s, %s, %s, %s, %s, %s, %s)""",
            (user_id, sid, e.lead, text, [f["id"]], c.ok, 100 + i, Jsonb(c.as_json())))

    for i, (ed, c) in enumerate(checked.education):
        conn.execute(
            """insert into education (user_id, institution, degree, meta, result, lines, sort, confirmed, provenance)
               values (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (user_id, ed.institution, ed.degree, ed.meta, ed.result, ed.lines, 100 + i, c.ok, Jsonb(c.as_json())))

    for s in checked.skills:
        conn.execute(
            """insert into facts (user_id, kind, text, source, confirmed_at, provenance)
               values (%s, 'skill', %s, 'upload', case when %s then now() end, %s)""",
            (user_id, s.text, s.check.ok, Jsonb(s.check.as_json())))
    for t in checked.other_facts:
        conn.execute(
            """insert into facts (user_id, kind, text, source, confirmed_at, provenance)
               values (%s, 'other', %s, 'upload', case when %s then now() end, %s)""",
            (user_id, t.text, t.check.ok, Jsonb(t.check.as_json())))


def usable_record(conn) -> tuple[list[dict], list[str], list[str], list[dict]]:
    """Usable items (with bullets), skills, award/role lines and education, for the model."""
    items = conn.execute("select id, key, kind, name, tagline, period, stack from items where confirmed order by sort").fetchall()
    bullets = conn.execute("select item_id, text from bullets where confirmed order by sort").fetchall()
    by_item: dict = {}
    for b in bullets:
        by_item.setdefault(b["item_id"], []).append(b["text"])
    rec = [{"key": i["key"], "kind": i["kind"], "name": i["name"], "tagline": i["tagline"], "period": i["period"],
            "stack": i["stack"], "bullets": by_item.get(i["id"], [])} for i in items]
    skills = [r["text"] for r in conn.execute(
        "select text from facts where kind = 'skill' and confirmed_at is not null order by created_at").fetchall()]
    entries = [((r["lead"] or "") + r["text"]).strip() for r in conn.execute(
        "select lead, text from entries where confirmed order by sort").fetchall()]
    edu = [dict(r) for r in conn.execute(
        "select institution, degree, meta, result from education where confirmed order by sort").fetchall()]
    return rec, skills, entries, edu


def propose_baselines(user_id: str, families: list[str], corpus: Corpus) -> tuple[list[dict], list[dict]]:
    with user_tx(user_id) as conn:
        items, skills, entries, edu = usable_record(conn)
        prof = conn.execute("""select name, career_stage, experience_years, experience_band, grad_date,
                                      batch_year from profiles""").fetchone() or {}
    years = float(prof["experience_years"]) if prof.get("experience_years") is not None else None
    person = {**{k: prof.get(k) for k in ("name", "career_stage", "experience_band", "grad_date", "batch_year")},
              "experience_years": years, "whole_years": int(years) if years else 0, "education": edu}
    out = llm.structured("onb_baselines", BaselineSet, stable=[prompts.BASELINES],
                         volatile=baseline_context(families, person, items, skills, entries),
                         effort="medium", max_tokens=10000, ctx=llm.CallContext(user_id=user_id))
    return validate_baselines(out, families, items, skills, corpus, prof.get("experience_band") or "entry",
                              extra_numbers_for(person))


def store_tracks(conn, user_id: str, tracks: list[dict], replace_all: bool) -> None:
    """Insert or replace baseline plans; with replace_all, families no longer targeted lose
    their track and baseline file. Call inside user_tx(user_id)."""
    keys = [t["key"] for t in tracks]
    if replace_all:
        conn.execute("""delete from resume_files where id in (select file_id from resumes
                        where is_baseline and not (track_key = any(%s)))""", (keys,))
        conn.execute("delete from tracks where not (key = any(%s))", (keys,))
    for t in tracks:
        conn.execute(
            """insert into tracks (user_id, key, label, title_line, summary, left_sections, skills, sort, approved,
                   role_family, fit, fit_why, gaps)
               values (%s, %s, %s, %s, %s, %s, %s, %s, true, %s, %s, %s, %s)
               on conflict (user_id, key) do update set label = excluded.label, title_line = excluded.title_line,
                   summary = excluded.summary, left_sections = excluded.left_sections, skills = excluded.skills,
                   sort = excluded.sort, approved = true, scale = null, role_family = excluded.role_family,
                   fit = excluded.fit, fit_why = excluded.fit_why, gaps = excluded.gaps""",
            (user_id, t["key"], t["label"], t["title_line"], t["summary"], Jsonb(t["left_sections"]), Jsonb(t["skills"]),
             t["sort"], t["role_family"], t["fit"], t["fit_why"], t["gaps"]))


def build_profile(user_id: str) -> dict:
    """The whole background build, with progress on profiles.build."""
    set_build(user_id, status="running", step="reading", error=None, started_at=_now())
    docs = _gather_sources(user_id)
    if not docs:
        raise ValueError("Nothing to read: upload a CV or write about your work")
    corpus = Corpus(docs)
    ex = extract_profile(docs, llm.CallContext(user_id=user_id))

    set_build(user_id, step="checking")
    checked = check_extraction(ex, corpus)
    years = counted_experience(checked.items)
    with user_tx(user_id) as conn:
        store_checked(conn, user_id, checked)
        stage = conn.execute("select career_stage from profiles").fetchone()["career_stage"]
        conn.execute("update profiles set experience_years = %s, experience_band = %s",
                     (years, band_for_years(years, stage)))
        families = conn.execute("select target_families from preferences").fetchone()["target_families"] or []
    evidence_check(user_id)
    if not families:
        raise ValueError("Pick at least one kind of role")

    set_build(user_id, step="writing", **checked.counts())
    tracks, suggestions = propose_baselines(user_id, families, corpus)
    with user_tx(user_id) as conn:
        store_tracks(conn, user_id, tracks, replace_all=True)

    set_build(user_id, step="rendering", suggestions=suggestions)
    rendered = render_baselines(user_id, [t["key"] for t in tracks], report=False)
    with user_tx(user_id) as conn:
        conn.execute("""update profiles set onboarding_step = case when onboarding_step = 'done' then 'done'
                        else 'review' end""")
    set_build(user_id, status="done", step="done", finished_at=_now(),
              failed_tracks={k: v["error"] for k, v in rendered.items() if not v.get("ok")})
    return {"tracks": [t["key"] for t in tracks], "suggestions": suggestions, **checked.counts()}


def render_baselines(user_id: str, keys: Optional[list[str]] = None, report: bool = True) -> dict:
    """Fit each baseline to one page and store its PDF. With report, progress goes to
    profiles.build (a re-render after edits)."""
    from .pipeline import resume as resume_mod

    if report:
        set_build(user_id, status="running", step="rendering", error=None)
    try:
        out = resume_mod.calibrate_baselines(user_id, keys)
    except Exception as e:
        if report:
            set_build(user_id, status="failed", error=str(e)[:500])
        raise
    if report:
        set_build(user_id, status="done", step="done", finished_at=_now(),
                  failed_tracks={k: v["error"] for k, v in out.items() if not v.get("ok")})
    return out


def add_family(user_id: str, family: str) -> dict:
    """A resume for one more kind of role, written and rendered like the others."""
    if family not in FAMILIES or family == "other":
        raise ValueError("Unknown kind of role")
    with user_tx(user_id) as conn:
        families = conn.execute("select target_families from preferences").fetchone()["target_families"] or []
        if family in families:
            return {"tracks": families}
        if len(families) >= MAX_FAMILIES:
            raise ValueError("Up to {} kinds of role: remove one first".format(MAX_FAMILIES))
        families = families + [family]
        conn.execute("update preferences set target_families = %s, role_types = %s",
                     (families, role_types(families)))
        corpus = stored_corpus(conn)
    set_build(user_id, status="running", step="writing", error=None)
    tracks, _ = propose_baselines(user_id, [family], corpus)
    tracks[0]["sort"] = len(families) - 1
    with user_tx(user_id) as conn:
        store_tracks(conn, user_id, tracks, replace_all=False)
    render_baselines(user_id, [family])
    return {"tracks": families}


def remove_family(user_id: str, family: str) -> list[str]:
    with user_tx(user_id) as conn:
        families = conn.execute("select target_families from preferences").fetchone()["target_families"] or []
        if family not in families:
            return families
        if len(families) == 1:
            raise ValueError("Keep at least one kind of role")
        families = [f for f in families if f != family]
        conn.execute("update preferences set target_families = %s, role_types = %s",
                     (families, role_types(families)))
        conn.execute("delete from resume_files where id in (select file_id from resumes where track_key = %s and is_baseline)",
                     (family,))
        conn.execute("delete from tracks where key = %s", (family,))
    return families


def finish(user_id: str) -> None:
    from .pipeline.run import queue_rematch
    with user_tx(user_id) as conn:
        conn.execute("update profiles set onboarding_step = 'done'")
    register_segments(user_id)
    queue_rematch(user_id)


def register_segments(user_id: str) -> list[tuple[str, str]]:
    """Make sure the pool has a segment for each (target family, band) this user needs, and
    recount demand. The collectors (phase 4) refresh segments by family."""
    with user_tx(user_id) as conn:
        families = conn.execute("select target_families from preferences").fetchone()["target_families"] or []
        band = conn.execute("select experience_band from profiles").fetchone()["experience_band"]
    if not band or not families:
        return []
    pairs = [(f, band) for f in families]
    with system_tx() as conn:
        for f, b in pairs:
            conn.execute("""insert into pool_segments (role_family, experience_band) values (%s, %s)
                            on conflict (role_family, experience_band) do update set status = 'active'""", (f, b))
        conn.execute(
            """update pool_segments s set demand = (
                   select count(*) from preferences p join profiles pr using (user_id)
                   where s.role_family = any(p.target_families) and pr.experience_band = s.experience_band
                     and pr.onboarding_step = 'done')
               where (s.role_family, s.experience_band) in (select unnest(%s::text[]), unnest(%s::text[]))""",
            ([f for f, _ in pairs], [b for _, b in pairs]))
    return pairs


def fail_build(user_id: str, message: str) -> None:
    set_build(user_id, status="failed", error=message[:500])


def evidence_check(user_id: str) -> list[dict]:
    """Which usable items or entries mention each skill. Not shown as a step any more: it
    orders skills and, later, weighs matches (a skill a project shows counts more)."""
    with user_tx(user_id) as conn:
        skills = conn.execute("select id::text, text from facts where kind = 'skill' order by text").fetchall()
        items = conn.execute("select id, key, name, coalesce(stack, '') as stack from items where confirmed").fetchall()
        bullets = conn.execute("select item_id, text from bullets where confirmed").fetchall()
        entries = conn.execute("select id::text, coalesce(lead, '') || text as text from entries where confirmed").fetchall()
        text_by_item: dict = {i["id"]: i["name"] + " " + i["stack"] for i in items}
        for b in bullets:
            text_by_item[b["item_id"]] = text_by_item.get(b["item_id"], "") + " " + b["text"]
        out = []
        for s in skills:
            rx = re.compile(r"(?<![\w]){}(?![\w])".format(re.escape(s["text"].lower())))
            backed = [i["key"] for i in items if rx.search(text_by_item.get(i["id"], "").lower())]
            backed += ["entry:" + e["id"] for e in entries if rx.search(e["text"].lower())]
            conn.execute("update facts set evidence_items = %s where id = %s", (backed, s["id"]))
            out.append({"fact_id": s["id"], "skill": s["text"], "backed_by": backed, "backed": bool(backed)})
    return out

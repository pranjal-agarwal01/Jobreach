"""
Two-column one-page CV builder. Port of the reference pipeline's build_resume.py
(schema v2), generalised to schema v3 and made safe to run concurrently: the active
type scale lives on a builder instance instead of a module global.

The format was measured off the original user's own CVs:

  A4 595.3 x 841.9pt, margins L30 R31 T41 B42
  One 1-row, 2-cell borderless table filling the page
    left  cell 340pt wide (329pt content, 11pt right pad)  -> name, summary, left sections
    right cell 195pt wide (184pt content, 11pt left pad)   -> contact ... right sections
    divider = left cell's right border, 0.75pt, #D6DCE5
  Accent #1F3864, secondary text #444444, Calibri throughout
  Section headings 10pt bold accent + accent bottom rule across the column

TYPE SCALE
  Every size and gap is a BASE value at scale 1.0. `scale` multiplies all of them
  together so proportions never drift. The page geometry does not scale. A track's
  scale is the largest at which it still renders as exactly one page, found by
  calibrate.py against a real renderer. Any content change can move it.

Design rules carried over (they are the format, not accidents):
  * The columns are a real Word table so ATS text extraction reads cell by cell.
    Never fake them with tab stops or text boxes.
  * Separators: U+00B7 between a title and its meta trailer, U+2013 for date
    ranges, U+2014 in prose. The no-em-dash rule applies to email, not to the CV.
  * ONE PAGE, always. Never ship a build that spills.
"""
from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .schema import BuildOptions, Item, LeftSection, ResumeData

# ---- page + column geometry, in points. Geometry does NOT scale. ----
PAGE_W, PAGE_H = 595.3, 841.9
M_LEFT, M_RIGHT, M_TOP, M_BOTTOM = 30.0, 31.0, 41.0, 42.0
L_CELL_W, R_CELL_W = 340.0, 195.0
GUTTER_PAD = 11.0
L_TEXT_W = L_CELL_W - GUTTER_PAD          # 329pt
R_TEXT_W = R_CELL_W - GUTTER_PAD          # 184pt
USABLE_H = PAGE_H - M_TOP - M_BOTTOM      # 758.9pt

ACCENT = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x44, 0x44, 0x44)
RULE_LIGHT = "D6DCE5"
ACCENT_HEX = "1F3864"

# ---- base type scale (scale 1.0 = the original measured CVs) ----
NAME_PT = 19.0
TITLE_PT = 9.5
SUMMARY_PT = 9.0
HEAD_PT = 10.0
PROJ_PT = 9.5
META_PT = 8.0
STACK_PT = 8.0
BODY_PT = 9.0        # left-column bullets
R_BODY_PT = 8.0      # right column
EDU_INST_PT = 8.5
EDU_META_PT = 7.5

# Calibri. Google Docs has no Calibri and substitutes a wider face, so a file that is
# one page in Word can preview as two in Google Drive. The server renders with
# Carlito, which is metric-compatible with Calibri. Tell users to view in Word.
FONT = "Calibri"

CHAR_W = 0.45        # Calibri mean advance / point size, calibrated on the original CVs

# The estimate is a GUIDE, NOT A GATE: it cannot know the renderer's real line
# breaking and does not even rank builds correctly near the one-page boundary. Only
# pages.count_pages decides whether a build is one page. The estimate is useful for
# telling which column is longer, and as a coarse alarm for a big overrun.
FIT_THRESHOLD = 1.25


@dataclass
class BuildResult:
    path: str
    track: str
    scale: float
    est_pages: float
    est_left: float
    est_right: float
    item_ids: list[str] = field(default_factory=list)
    bullet_ids: list[str] = field(default_factory=list)
    entry_ids: list[str] = field(default_factory=list)


def _twips(pt: float) -> str:
    return str(int(round(pt * 20)))


def _bottom_rule(p, color=ACCENT_HEX, sz="6"):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def _cell_margins(cell, left=0.0, right=0.0, top=0.0, bottom=0.0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), _twips(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _cell_borders(cell, right_rule=None):
    """All borders off; optionally a single coloured rule on the right edge."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for tag in ("top", "start", "bottom", "end"):
        el = OxmlElement("w:" + tag)
        if tag == "end" and right_rule:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), right_rule)
        else:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
            el.set(qn("w:space"), "0")
        borders.append(el)
    tcPr.append(borders)


def _cell_width(cell, pt):
    tcPr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), _twips(pt))
    w.set(qn("w:type"), "dxa")
    tcPr.append(w)


def _fixed_layout(table, total_pt):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), _twips(total_pt))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)


def _tail_paragraph(doc):
    """End the document with a 1pt empty paragraph.

    A document cannot end on a table, so without this Word inserts a body-size empty
    paragraph after it. When the table runs close to the bottom margin, that paragraph
    lands on a blank page 2, which Word counts as a spill and the recipient sees.
    LibreOffice does not do this, which made the two renderers disagree at exactly the
    one-page boundary. A 1pt paragraph fits in any leftover space in both.
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    rPr = OxmlElement("w:rPr")   # paragraph-mark size: an empty paragraph's line height
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rPr.append(sz)
    p._p.get_or_add_pPr().append(rPr)


def _para(cell, first):
    """Reuse the empty paragraph Word puts in a fresh cell, then append."""
    return cell.paragraphs[0] if first else cell.add_paragraph()


class _Builder:
    def __init__(self, scale: float):
        self.scale = float(scale)

    def s(self, v: float) -> float:
        """Scale a point value. Rounded to 1/4pt; Word stores half-points."""
        return round(v * self.scale * 4) / 4.0

    # ---------------------------------------------------------------- primitives

    def spacing(self, p, before=0.0, after=0.0, line=1.0):
        pf = p.paragraph_format
        pf.space_before = Pt(self.s(before))
        pf.space_after = Pt(self.s(after))
        pf.line_spacing = line
        return p

    def run(self, p, text, size, bold=False, italic=False, color=None):
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(self.s(size))
        r.font.name = FONT
        if color is not None:
            r.font.color.rgb = color
        return r

    def hyperlink(self, p, text, url, size, color=ACCENT, bold=False, underline=False):
        """python-docx has no hyperlink API; build the w:hyperlink element by hand."""
        r_id = p.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        rPr.append(fonts)
        if bold:
            rPr.append(OxmlElement("w:b"))
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "%02X%02X%02X" % (color[0], color[1], color[2]))
        rPr.append(col)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(round(self.s(size) * 2))))
        rPr.append(sz)
        if underline:
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rPr.append(u)

        run.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        link.append(run)
        p._p.append(link)
        return link

    def heading(self, cell, text, first=False, before=9):
        p = _para(cell, first)
        self.spacing(p, before=0 if first else before, after=3.5)
        self.run(p, text.upper(), HEAD_PT, bold=True, color=ACCENT)
        _bottom_rule(p)
        return p

    def bullet(self, cell, text, size, indent, after=2.5, lead=None, glyph_pt=None):
        """Hanging-indent bullet. `lead` renders bold ahead of `text`."""
        p = cell.add_paragraph()
        self.spacing(p, after=after, line=1.0)
        pf = p.paragraph_format
        pf.left_indent = Pt(self.s(indent))
        pf.first_line_indent = Pt(-self.s(indent))
        pf.tab_stops.add_tab_stop(Pt(self.s(indent)))
        self.run(p, "•", glyph_pt if glyph_pt else size - 1)
        if lead:
            self.run(p, "\t" + lead, size, bold=True)
            self.run(p, text, size)
        else:
            self.run(p, "\t" + text, size)
        return p

    # ---------------------------------------------------------------- columns

    def left_column(self, cell, data: ResumeData, track, sections: list[tuple[str, list[Item]]],
                    drop: set[str], result: BuildResult):
        c = data.contact

        p = _para(cell, first=True)
        self.spacing(p, after=1)
        self.run(p, c.name, NAME_PT, bold=True, color=ACCENT)

        p = cell.add_paragraph()
        self.spacing(p, after=4)
        self.run(p, track.title_line, TITLE_PT, color=GRAY)

        if track.summary:
            p = cell.add_paragraph()
            self.spacing(p, after=0, line=1.0)
            self.run(p, track.summary, SUMMARY_PT)

        for heading, items in sections:
            if not items:
                continue
            self.heading(cell, heading, before=9)
            for i, item in enumerate(items):
                result.item_ids.append(item.id)
                p = cell.add_paragraph()
                self.spacing(p, before=0 if i == 0 else 7, after=1)
                self.run(p, item.name, PROJ_PT, bold=True)
                if item.tagline:
                    self.run(p, " — " + item.tagline, PROJ_PT, bold=True)
                for link in item.links:
                    self.run(p, " | ", META_PT, color=GRAY)
                    self.hyperlink(p, link.text, link.url, META_PT, underline=True)
                if item.period:
                    self.run(p, " · " + item.period, META_PT, color=GRAY)

                if item.stack:
                    p = cell.add_paragraph()
                    self.spacing(p, after=2.5, line=1.0)
                    self.run(p, item.stack_label + ": ", STACK_PT, bold=True, color=GRAY)
                    self.run(p, item.stack, STACK_PT, color=GRAY)

                for b in item.bullets:
                    if b.id in drop:
                        continue
                    result.bullet_ids.append(b.id)
                    self.bullet(cell, b.text, BODY_PT, 8.2)

    def right_column(self, cell, data: ResumeData, track, drop: set[str], result: BuildResult):
        c = data.contact
        first = True

        contact_lines = [line for line in (c.location, c.phone) if line]
        if contact_lines or c.email:
            self.heading(cell, "Contact", first=first)
            first = False
            for line in contact_lines:
                p = cell.add_paragraph()
                self.spacing(p, after=1.5, line=1.0)
                self.run(p, line, R_BODY_PT)
            if c.email:
                p = cell.add_paragraph()
                self.spacing(p, after=1.5, line=1.0)
                self.hyperlink(p, c.email, "mailto:" + c.email, R_BODY_PT, underline=True)

        if c.social:
            self.heading(cell, "Social", first=first)
            first = False
            for s in c.social:
                p = cell.add_paragraph()
                self.spacing(p, after=1.5, line=1.0)
                self.hyperlink(p, s.text, s.url, R_BODY_PT, underline=True)

        if track.skills:
            self.heading(cell, "Technical Skills", first=first)
            first = False
            for group in track.skills:
                p = cell.add_paragraph()
                self.spacing(p, after=2.5, line=1.0)
                self.run(p, group.label + ": ", R_BODY_PT, bold=True)
                self.run(p, group.items, R_BODY_PT)

        if data.education:
            self.heading(cell, "Education", first=first)
            first = False
            for i, ed in enumerate(data.education):
                p = cell.add_paragraph()
                self.spacing(p, before=0 if i == 0 else 5, after=0.5, line=1.0)
                self.run(p, ed.institution, EDU_INST_PT, bold=True)
                if ed.degree:
                    p = cell.add_paragraph()
                    self.spacing(p, after=0.5, line=1.0)
                    self.run(p, ed.degree, R_BODY_PT)
                if ed.meta:
                    p = cell.add_paragraph()
                    self.spacing(p, after=0.5, line=1.0)
                    self.run(p, ed.meta, EDU_META_PT, color=GRAY)
                if ed.result:
                    p = cell.add_paragraph()
                    self.spacing(p, after=0.5, line=1.0)
                    self.run(p, ed.result, R_BODY_PT, bold=True)
                for line in ed.lines:
                    p = cell.add_paragraph()
                    self.spacing(p, after=0.5, line=1.0)
                    self.run(p, line, R_BODY_PT)

        for sec in data.sections:
            entries = [e for e in sec.entries
                       if e.id not in drop and (e.tracks is None or track.key in e.tracks)]
            if not entries:
                continue
            self.heading(cell, sec.heading, first=first)
            first = False
            after = 2.0 if sec.style == "list" else 2.5
            for e in entries:
                result.entry_ids.append(e.id)
                self.bullet(cell, e.text, R_BODY_PT, 7.4, after=after,
                            lead=e.lead, glyph_pt=R_BODY_PT)

    # ---------------------------------------------------------------- estimate

    def column_height(self, cell, width_pt):
        """Simulate wrapping to estimate a column's rendered height in points."""
        total = 0.0
        for p in cell.paragraphs:
            txt = p.text
            pf = p.paragraph_format
            sizes = [r.font.size.pt for r in p.runs if r.font.size]
            size = max(sizes) if sizes else self.s(BODY_PT)
            if not txt.strip():
                total += (pf.space_before.pt if pf.space_before else 0)
                total += (pf.space_after.pt if pf.space_after else 0)
                continue
            indent = pf.left_indent.pt if pf.left_indent else 0.0
            avail = width_pt - indent
            cpl = max(1, int(avail / (size * CHAR_W)))
            lines = max(1, -(-len(txt) // cpl))
            mult = pf.line_spacing if isinstance(pf.line_spacing, float) else 1.0
            total += lines * size * 1.22 * mult
            total += (pf.space_before.pt if pf.space_before else 0)
            total += (pf.space_after.pt if pf.space_after else 0)
            if p._p.find(qn("w:pPr") + "/" + qn("w:pBdr")) is not None:
                total += 2.0
        return total


def resolve_sections(data: ResumeData, opts: BuildOptions) -> list[tuple[str, list[Item]]]:
    track = data.tracks[opts.track]
    left: list[LeftSection] = opts.left_sections or track.left_sections
    out = []
    for sec in left:
        missing = [i for i in sec.item_ids if i not in data.items]
        if missing:
            raise ValueError("unknown item id(s): {}".format(", ".join(missing)))
        out.append((sec.heading, [data.items[i] for i in sec.item_ids if i not in opts.drop]))
    return out


def build(data: ResumeData, opts: BuildOptions, out_path: str,
          scale: Optional[float] = None) -> BuildResult:
    """Write a one-table two-column .docx. Does not check pages; see pages.py."""
    if opts.track not in data.tracks:
        raise ValueError("unknown track {!r}".format(opts.track))
    track = data.tracks[opts.track]
    if scale is None:
        scale = opts.scale if opts.scale is not None else (track.scale or 1.0)
    b = _Builder(scale)
    result = BuildResult(path=out_path, track=opts.track, scale=b.scale,
                         est_pages=0.0, est_left=0.0, est_right=0.0)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(b.s(BODY_PT))
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    for s in doc.sections:
        s.page_width, s.page_height = Pt(PAGE_W), Pt(PAGE_H)
        s.left_margin, s.right_margin = Pt(M_LEFT), Pt(M_RIGHT)
        s.top_margin, s.bottom_margin = Pt(M_TOP), Pt(M_BOTTOM)

    sections = resolve_sections(data, opts)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _fixed_layout(table, L_CELL_W + R_CELL_W)
    # The table grid must carry the real column widths. python-docx initialises it to
    # two equal halves (267pt each). Word lays out by the cell widths and hides the
    # mismatch; LibreOffice (and likely Google Docs) lays out by the grid, squeezing
    # the left column to 267pt and pushing it onto page 2. The reference pipeline's
    # files all carry that mismatch.
    table.columns[0].width = Pt(L_CELL_W)
    table.columns[1].width = Pt(R_CELL_W)

    left, right = table.rows[0].cells
    _cell_width(left, L_CELL_W)
    _cell_width(right, R_CELL_W)
    _cell_margins(left, left=0, right=GUTTER_PAD, top=0, bottom=0)
    _cell_margins(right, left=GUTTER_PAD, right=0, top=0, bottom=0)
    _cell_borders(left, right_rule=RULE_LIGHT)
    _cell_borders(right)

    b.left_column(left, data, track, sections, opts.drop, result)
    b.right_column(right, data, track, opts.drop, result)
    _tail_paragraph(doc)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)

    lh = b.column_height(left, L_TEXT_W) / USABLE_H
    rh = b.column_height(right, R_TEXT_W) / USABLE_H
    result.est_pages, result.est_left, result.est_right = max(lh, rh), lh, rh
    return result

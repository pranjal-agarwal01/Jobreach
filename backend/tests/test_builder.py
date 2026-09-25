import re
import zipfile

import pytest
from docx import Document

from resume_engine.ats_score import all_paragraphs, is_two_column_layout
from resume_engine.builder import build
from resume_engine.schema import BuildOptions, LeftSection

from .conftest import PROFILES, load_profile

TRACK_CASES = [(p, t) for p in PROFILES for t in load_profile(p).tracks]


@pytest.mark.parametrize("path,track", TRACK_CASES, ids=lambda v: getattr(v, "stem", v))
def test_every_profile_track_builds_as_one_two_cell_table(path, track, tmp_path):
    data = load_profile(path)
    out = tmp_path / "resume.docx"
    result = build(data, BuildOptions(track=track), str(out))
    doc = Document(str(out))
    assert is_two_column_layout(doc)
    text = "\n".join(p.text for p in all_paragraphs(doc))
    assert data.contact.name in text
    # Every bullet the builder reports is on the page, and nothing else from the pool is.
    for item in data.items.values():
        for b in item.bullets:
            assert (b.text in text) == (b.id in result.bullet_ids)
    assert result.est_pages < 1.25, "estimate says far over one page"


def test_track_filters_entries_and_sections(tmp_path):
    data = load_profile(next(p for p in PROFILES if "rohan" in p.stem))
    r = build(data, BuildOptions(track="backend"), str(tmp_path / "b.docx"))
    assert "aw3" in r.entry_ids and "aw2" not in r.entry_ids   # aws only on backend/fullstack
    assert "r2" not in r.entry_ids                            # open source: not backend
    assert r.item_ids == ["pixelforge", "tutorlink", "splitsy"]


def test_drop_and_reorder_select_without_adding(tmp_path):
    data = load_profile(next(p for p in PROFILES if "diya" in p.stem))
    opts = BuildOptions(track="ml", drop={"churn.b4", "c2"},
                        left_sections=[LeftSection(heading="Projects", item_ids=["churn", "reviewsent"])])
    r = build(data, opts, str(tmp_path / "t.docx"))
    assert r.item_ids == ["churn", "reviewsent"]
    assert "churn.b4" not in r.bullet_ids and "churn.b1" in r.bullet_ids
    assert "c2" not in r.entry_ids
    text = "\n".join(p.text for p in all_paragraphs(Document(str(tmp_path / "t.docx"))))
    assert "EXPERIENCE" not in text   # override replaced the track's sections


def test_experience_heading_and_missing_phone(tmp_path):
    diya = load_profile(next(p for p in PROFILES if "diya" in p.stem))
    build(diya, BuildOptions(track="data"), str(tmp_path / "d.docx"))
    text = "\n".join(p.text for p in all_paragraphs(Document(str(tmp_path / "d.docx"))))
    assert text.index("EXPERIENCE") < text.index("PROJECTS")
    assert "Tools: SQL, Python" in text

    kabir = load_profile(next(p for p in PROFILES if "kabir" in p.stem))
    assert kabir.contact.phone is None
    build(kabir, BuildOptions(track="embedded"), str(tmp_path / "k.docx"))


def test_table_grid_matches_cell_widths(tmp_path):
    data = load_profile(PROFILES[0])
    build(data, BuildOptions(track="sde"), str(tmp_path / "g.docx"))
    xml = zipfile.ZipFile(tmp_path / "g.docx").read("word/document.xml").decode()
    assert re.findall(r'<w:gridCol w:w="(\d+)"/>', xml) == re.findall(r'<w:tcW w:w="(\d+)"', xml)


def test_unknown_track_rejected(tmp_path):
    with pytest.raises(ValueError):
        build(load_profile(PROFILES[0]), BuildOptions(track="nope"), str(tmp_path / "x.docx"))
